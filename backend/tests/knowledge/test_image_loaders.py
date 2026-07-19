"""Unit tests for image ingestion — caption + OCR + embedded images (ITEM 2 Slice C).

Uses a **stub** vision captioner / OCR (no real model, no tesseract binary) injected
through :class:`LoadContext`, so the tests are deterministic and offline. Covers:
the standalone image loader, the caption/OCR combination, the honest empty-image
failure, embedded images in a docx, and the v1 image message format the real
captioner emits.
"""

from __future__ import annotations

import io

import pytest
from langchain_core.messages import AIMessage

from app.knowledge.image_processing import make_vision_captioner
from app.knowledge.loaders import (
    EmptyDocument,
    LoadContext,
    image_text_from,
    load_file_bytes,
)


def _png(color: tuple[int, int, int] = (200, 30, 30)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), color).save(buf, format="PNG")
    return buf.getvalue()


def _caption_ctx(text: str = "A red square test image") -> LoadContext:
    return LoadContext(caption_image=lambda data, mime: text)


def test_image_text_from_combines_caption_and_ocr() -> None:
    ctx = LoadContext(
        caption_image=lambda d, m: "a diagram",
        ocr_image=lambda d: "INVOICE 2026",
    )
    out = image_text_from(b"x", "image/png", ctx)
    assert "a diagram" in out
    assert "INVOICE 2026" in out


def test_image_text_from_empty_when_no_services() -> None:
    assert image_text_from(b"x", "image/png", LoadContext()) == ""


def test_standalone_image_loads_with_caption() -> None:
    docs = load_file_bytes(_png(), "photo.png", _caption_ctx("A red square"))
    assert docs[0].page_content == "A red square"
    assert docs[0].metadata["kind"] == "image"


def test_standalone_image_without_services_fails_clearly() -> None:
    # No vision + no OCR → no text → EmptyDocument (worker reports a precise reason),
    # never a silent empty ingest.
    with pytest.raises(EmptyDocument):
        load_file_bytes(_png(), "photo.png", LoadContext())


def test_jpg_and_webp_are_registered_image_formats() -> None:
    from app.knowledge.loaders import supported_file_formats

    assert {"png", "jpg", "jpeg", "webp", "tiff"} <= supported_file_formats()


def test_docx_embedded_image_is_captioned() -> None:
    from docx import Document as DocxDocument

    d = DocxDocument()
    d.add_paragraph("Report body text.")
    d.add_picture(io.BytesIO(_png()))
    buf = io.BytesIO()
    d.save(buf)

    docs = load_file_bytes(buf.getvalue(), "report.docx", _caption_ctx("Embedded chart"))
    # One text document + one image-derived document.
    assert any("Report body text." in doc.page_content for doc in docs)
    image_docs = [doc for doc in docs if doc.metadata.get("image")]
    assert image_docs and image_docs[0].page_content == "Embedded chart"


def test_vision_captioner_emits_v1_image_block() -> None:
    captured: dict[str, object] = {}

    class _FakeModel:
        def invoke(self, messages: list[object]) -> AIMessage:
            captured["messages"] = messages
            return AIMessage(content="A bar chart of revenue")

    caption = make_vision_captioner(_FakeModel())  # type: ignore[arg-type]
    result = caption(_png(), "image/png")
    assert result == "A bar chart of revenue"

    # The image was sent as the langchain v1 standard content block.
    human = captured["messages"][0]  # type: ignore[index]
    blocks = human.content
    image_block = next(b for b in blocks if isinstance(b, dict) and b.get("type") == "image")
    assert image_block["mime_type"] == "image/png"
    assert isinstance(image_block["base64"], str) and image_block["base64"]
