"""Unit tests for the office/media byte generators (ARTIFACTS §5, Slice 4).

Pure functions, no DB: each generator produces a valid file (verified by magic bytes
and re-opening with the source library) and rejects malformed specs with ValueError
(which the tools turn into a refusal the agent can correct).
"""

from __future__ import annotations

import io

import pytest

from app.artifacts.generators import (
    build_chart,
    build_docx,
    build_image,
    build_pdf,
    build_pptx,
    build_xlsx,
)


def test_build_docx_is_valid_and_contains_text() -> None:
    from docx import Document

    data = build_docx(title="Report", sections=[{"heading": "Intro", "body": "p1\n\np2"}])
    assert data[:2] == b"PK"  # OOXML zip
    texts = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert "Report" in texts and "Intro" in texts and "p1" in texts


def test_build_xlsx_is_valid_and_keeps_cells() -> None:
    from openpyxl import load_workbook

    data = build_xlsx(sheets=[{"name": "Q1", "rows": [["Item", "Cost"], ["Widget", 10]]}])
    assert data[:2] == b"PK"
    workbook = load_workbook(io.BytesIO(data))
    assert workbook["Q1"]["A1"].value == "Item"
    assert workbook["Q1"]["B2"].value == 10


def test_build_pptx_is_valid_with_one_slide() -> None:
    from pptx import Presentation

    data = build_pptx(slides=[{"title": "Plan", "bullets": ["a", "b"]}])
    assert data[:2] == b"PK"
    assert len(Presentation(io.BytesIO(data)).slides) == 1


def test_build_pdf_is_valid() -> None:
    data = build_pdf(content_md="# Title\n\nSome body text.")
    assert data[:5] == b"%PDF-"


def test_build_chart_is_a_png() -> None:
    data = build_chart(spec={"type": "bar", "labels": ["a", "b"], "values": [1, 2], "title": "t"})
    assert data[:8] == b"\x89PNG\r\n\x1a\n"


def test_build_image_is_a_png() -> None:
    data = build_image(spec={"width": 120, "height": 60, "text": "hi"})
    assert data[:8] == b"\x89PNG\r\n\x1a\n"


def test_build_xlsx_requires_a_sheet() -> None:
    with pytest.raises(ValueError):
        build_xlsx(sheets=[])


def test_build_chart_requires_labels() -> None:
    with pytest.raises(ValueError):
        build_chart(spec={"type": "bar", "labels": [], "values": []})


def test_build_chart_rejects_series_length_mismatch() -> None:
    with pytest.raises(ValueError):
        build_chart(spec={"type": "line", "labels": ["a", "b"], "values": [1]})


def test_build_pdf_requires_content() -> None:
    with pytest.raises(ValueError):
        build_pdf(content_md="   ")
